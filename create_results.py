import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.formatting.rule import ColorScaleRule


# ---------------------------------------------------------------------------
# Wczytanie i połączenie wyników
# ---------------------------------------------------------------------------

df_classic = pd.read_csv("wyniki_modele_klasyczne.csv")
df_fnn = pd.read_csv("wyniki_fnn.csv")

df = pd.concat([df_classic, df_fnn], ignore_index=True)

# Kolejność kolumn (na wypadek gdyby w plikach była inna)
column_order = [
    "selection_method", "n_features", "model", "selected_features",
    "TN", "FP", "FN", "TP",
    "accuracy", "precision", "recall", "f1", "roc_auc", "pr_auc",
    "cost_5", "cost_10", "cost_20",
]
df = df[[c for c in column_order if c in df.columns]]

# Sortowanie: metoda selekcji, liczba cech, model
df = df.sort_values(by=["selection_method", "n_features", "model"]).reset_index(drop=True)

column_labels = {
    "selection_method": "Metoda selekcji",
    "n_features": "Liczba cech",
    "model": "Model",
    "selected_features": "Wybrane cechy",
    "TN": "TN",
    "FP": "FP",
    "FN": "FN",
    "TP": "TP",
    "accuracy": "Accuracy",
    "precision": "Precision",
    "recall": "Recall",
    "f1": "F1-Score",
    "roc_auc": "ROC-AUC",
    "pr_auc": "PR-AUC",
    "cost_5": "Koszt (FN=5)",
    "cost_10": "Koszt (FN=10)",
    "cost_20": "Koszt (FN=20)",
}
headers = [column_labels[c] for c in df.columns]


# ---------------------------------------------------------------------------
# Budowa arkusza Excel
# ---------------------------------------------------------------------------

FONT_NAME = "Arial"

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Wyniki"

header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
header_font = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=11)
normal_font = Font(name=FONT_NAME, size=11)
center = Alignment(horizontal="center", vertical="center")
thin = Side(style="thin", color="B7B7B7")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

model_fills = {
    "LogisticRegression": PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid"),
    "RandomForest":       PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"),
    "XGBoost":            PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),
    "SVM":                PatternFill(start_color="E4DFEC", end_color="E4DFEC", fill_type="solid"),
    "FNN":                PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"),
}

# Nagłówek
ws.append(headers)
for col_idx in range(1, len(headers) + 1):
    c = ws.cell(row=1, column=col_idx)
    c.font = header_font
    c.fill = header_fill
    c.alignment = center
    c.border = border

# Kolumny procentowe/liczbowe wg nazwy
pct_cols = {"accuracy", "precision", "recall", "f1", "roc_auc", "pr_auc"}
int_cols = {"cost_5", "cost_10", "cost_20"}

for row in dataframe_to_rows(df, index=False, header=False):
    ws.append(row)

last_row = ws.max_row

for row_idx in range(2, last_row + 1):
    model_name = ws.cell(row=row_idx, column=df.columns.get_loc("model") + 1).value
    fill = model_fills.get(model_name)

    for col_idx, col_name in enumerate(df.columns, start=1):
        c = ws.cell(row=row_idx, column=col_idx)
        c.font = normal_font
        c.border = border
        c.alignment = center
        if fill:
            c.fill = fill
        if col_name == "selected_features":
            c.alignment = Alignment(horizontal="left", vertical="center")
        if col_name in pct_cols:
            c.number_format = "0.0000"
        if col_name in int_cols:
            c.number_format = "#,##0"

# Szerokości kolumn
default_widths = {
    "selection_method": 18, "n_features": 12, "model": 20, "selected_features": 45,
    "TN": 9, "FP": 9, "FN": 9, "TP": 9,
    "accuracy": 11, "precision": 11, "recall": 10, "f1": 11,
    "roc_auc": 10, "pr_auc": 10,
    "cost_5": 13, "cost_10": 14, "cost_20": 14,
}
for i, col_name in enumerate(df.columns, start=1):
    ws.column_dimensions[get_column_letter(i)].width = default_widths.get(col_name, 14)

ws.freeze_panes = "A2"
ws.row_dimensions[1].height = 22

# Skala kolorów: F1-Score (im wyżej tym lepiej) i Koszt FN=20 (im niżej tym lepiej)
if "f1" in df.columns:
    f1_col = get_column_letter(df.columns.get_loc("f1") + 1)
    ws.conditional_formatting.add(
        f"{f1_col}2:{f1_col}{last_row}",
        ColorScaleRule(
            start_type="min", start_color="F8696B",
            mid_type="percentile", mid_value=50, mid_color="FFEB84",
            end_type="max", end_color="63BE7B",
        ),
    )

if "cost_20" in df.columns:
    cost_col = get_column_letter(df.columns.get_loc("cost_20") + 1)
    ws.conditional_formatting.add(
        f"{cost_col}2:{cost_col}{last_row}",
        ColorScaleRule(
            start_type="min", start_color="63BE7B",
            mid_type="percentile", mid_value=50, mid_color="FFEB84",
            end_type="max", end_color="F8696B",
        ),
    )

# ---------------------------------------------------------------------------
# Druga zakładka: legenda kolorów modeli
# ---------------------------------------------------------------------------

ws2 = wb.create_sheet("Legenda")
ws2["A1"] = "Model"
ws2["B1"] = "Kolor w tabeli"
ws2["A1"].font = header_font
ws2["B1"].font = header_font
ws2["A1"].fill = header_fill
ws2["B1"].fill = header_fill

models_present = df["model"].unique() if "model" in df.columns else []
r = 2
for model_name in models_present:
    fill = model_fills.get(model_name)
    ws2.cell(row=r, column=1, value=model_name).font = normal_font
    if fill:
        ws2.cell(row=r, column=2).fill = fill
    r += 1

ws2.column_dimensions["A"].width = 22
ws2.column_dimensions["B"].width = 18

wb.save("wyniki.xlsx")
print(f"Zapisano wyniki.xlsx ({len(df)} wierszy)")